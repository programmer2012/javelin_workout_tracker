#!/usr/bin/env python3
"""Import a coach's Excel training plan into Firestore.

Parses the Kennzahlen / Trainingsplan / AT-Split / Sprint & Sprung sheets and
writes the result as one document in `plans/{planId}`, which the app reads for
any athlete the plan is assigned to.

Usage
-----
    pip install openpyxl firebase-admin

    python tools/import_plan.py PLAN.xlsx \\
        --key serviceAccount.json \\
        --name "TP Speer Tom/Alex 2026/27"

    # see what would be written, without touching Firestore
    python tools/import_plan.py PLAN.xlsx --dry-run

Existing assignments are preserved on re-import, so correcting a typo does not
silently drop every athlete off the plan. Use --plan-id to overwrite a
specific plan; otherwise one is derived from the season.

The service-account key comes from the Firebase console under
Project settings -> Service accounts -> Generate new private key. Keep it out
of version control: tools/*.json is gitignored.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys

try:
    import openpyxl
except ImportError:  # pragma: no cover - a setup problem, not a bug
    sys.exit("openpyxl is missing:  pip install openpyxl")

DAYS = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]


# --- helpers ---------------------------------------------------------------

def text(value) -> str:
    return "" if value is None else str(value).strip()


def num(value):
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    m = re.search(r"\d+(?:[.,]\d+)?", str(value))
    return float(m.group(0).replace(",", ".")) if m else None


def week_number(value):
    """The calendar week in column A, or None if this is not a data row.

    Only a real numeric cell counts. Both sheets end in prose footnotes, and
    one reads "Wo 53 fehlt in der Vorlage ... Summen ohne Wo 53" - scraping a
    number out of that produced a second week 53 that overwrote the real one
    with zeros.
    """
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return None
    week = int(value)
    return week if 1 <= week <= 53 else None


# --- sheet parsers ---------------------------------------------------------

def parse_weeks(wb) -> list[dict]:
    weeks: dict[int, dict] = {}

    ws = wb["Kennzahlen"]
    for row in ws.iter_rows(min_row=5, values_only=True):
        week = week_number(row[0])
        if week is None:
            continue
        weeks[week] = {
            "week": week,
            "phase": text(row[5]),
            "sollPercent": {
                "at": num(row[1]) or 0,
                "kraft": num(row[2]) or 0,
                "spezKt": num(row[3]) or 0,
                "wuerfe": num(row[4]) or 0,
            },
            "soll": {
                "at": num(row[7]) or 0,
                "kraft": num(row[8]) or 0,
                "spezKt": num(row[9]) or 0,
                "wuerfe": num(row[10]) or 0,
            },
            "note": text(row[11]),
            "monday": None,
            "days": {},
        }

    ws = wb["Trainingsplan"]
    for row in ws.iter_rows(min_row=4, values_only=True):
        week = week_number(row[0])
        if week is None:
            continue
        entry = weeks.setdefault(week, {
            "week": week, "phase": "", "soll": {}, "sollPercent": {},
            "note": "", "monday": None, "days": {},
        })

        monday = row[1]
        if hasattr(monday, "year"):
            entry["monday"] = (
                f"{monday.year:04d}-{monday.month:02d}-{monday.day:02d}"
            )
        else:
            m = re.match(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", text(monday))
            if m:
                entry["monday"] = (
                    f"{int(m.group(3)):04d}-{int(m.group(2)):02d}"
                    f"-{int(m.group(1)):02d}"
                )

        if not entry["phase"]:
            entry["phase"] = text(row[2])
        for i, day in enumerate(DAYS):
            value = text(row[3 + i])
            if value:
                entry["days"][day] = value

    # The season starts in week 44 and wraps past 53 into the new year.
    return sorted(weeks.values(),
                  key=lambda w: (0 if w["week"] >= 44 else 1, w["week"]))


def parse_split(wb) -> dict:
    if "AT-Split" not in wb.sheetnames:
        return {"intro": [], "blocks": [], "progression": []}

    ws = wb["AT-Split"]
    intro, blocks, progression = [], [], []
    current, columns, mode = None, None, None

    for row in ws.iter_rows(min_row=1, values_only=True):
        cells = [text(c) for c in row]
        if not any(cells):
            continue
        first = cells[0]

        if first == "Nr":
            columns, mode = cells, "exercises"
            continue
        if first == "Woche":
            mode = "progression"
            continue

        if mode == "exercises" and re.fullmatch(r"\d+(\.0)?", first):
            if current is None:
                continue
            entry = {"nr": int(float(first)), "name": cells[1],
                     "equipment": cells[2]}
            third = cells[3] if len(cells) > 3 else ""
            if columns and len(columns) > 3 and "Typ" in columns[3]:
                entry["type"] = third
            elif third:
                entry["sets"] = third
            hint = cells[4] if len(cells) > 4 else ""
            if hint and hint != "–":
                entry["hint"] = hint
            current["exercises"].append(entry)
            continue

        if mode == "progression" and re.match(r"\d+", first):
            m = re.match(r"(\d+)\s*(.*)", first)
            progression.append({
                "week": int(m.group(1)), "label": m.group(2).strip(),
                "main": cells[1], "secondary": cells[2],
                "isolation": cells[3], "rir": cells[4], "focus": cells[5],
                "kraftReps": float(cells[6]) if cells[6] else 0,
                "medball": float(cells[7]) if cells[7] else 0,
                "atTotal": float(cells[8]) if cells[8] else 0,
                "soll": float(cells[9]) if cells[9] else 0,
            })
            continue

        if not any(cells[1:]):
            if first.startswith("•"):
                (current["notes"] if current else intro).append(
                    first.lstrip("• ").strip())
            else:
                current = {"title": first, "exercises": [], "notes": []}
                blocks.append(current)
            continue

        if current is None:
            intro.append(first)

    kept = []
    for block in blocks:
        if not block["exercises"] and not block["notes"]:
            intro.insert(0, block["title"])
        else:
            kept.append(block)

    return {"intro": intro, "blocks": kept, "progression": progression}


def parse_sprint(wb, season_order: list[int]) -> dict:
    if "Sprint & Sprung" not in wb.sheetnames:
        return {"blocks": [], "rules": []}

    ws = wb["Sprint & Sprung"]
    rows, rules, in_rules = [], [], False

    for row in ws.iter_rows(min_row=1, values_only=True):
        cells = [text(c) for c in row]
        if not any(cells):
            continue
        first = cells[0]

        if first == "Regeln":
            in_rules = True
            continue
        if in_rules:
            rules.append(first.lstrip("• ").strip())
            continue
        if first == "Wochen":
            continue

        m = re.fullmatch(r"(\d+)\s*[–\-]\s*(\d+)", first)
        if not m:
            continue
        start, end = int(m.group(1)), int(m.group(2))
        try:
            i, j = season_order.index(start), season_order.index(end)
        except ValueError:
            continue
        rows.append({
            "label": first, "weeks": season_order[i:j + 1],
            "goal": cells[1], "contents": cells[2], "volume": cells[3],
            "placement": cells[4] if len(cells) > 4 else "",
        })

    return {"blocks": rows, "rules": rules}


# --- checks ----------------------------------------------------------------

def check(plan: dict) -> list[str]:
    """Problems worth refusing to upload over."""
    problems = []
    weeks = plan["weeks"]

    if not weeks:
        problems.append("no weeks parsed")
        return problems

    for week in weeks:
        if not week["monday"]:
            problems.append(f"week {week['week']} has no Monday date")
        if not week["phase"]:
            problems.append(f"week {week['week']} has no phase")

    order = [w["week"] for w in weeks]
    if len(set(order)) != len(order):
        problems.append("a week number appears twice")

    covered = {w for b in plan["sprint"]["blocks"] for w in b["weeks"]}
    uncovered = [w for w in order if w not in covered]
    if plan["sprint"]["blocks"] and uncovered:
        problems.append(f"weeks with no sprint block: {uncovered}")

    return problems


# --- main ------------------------------------------------------------------

def build_plan(path: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    for required in ("Kennzahlen", "Trainingsplan"):
        if required not in wb.sheetnames:
            sys.exit(f"{path} has no '{required}' sheet "
                     f"(found: {', '.join(wb.sheetnames)})")

    weeks = parse_weeks(wb)
    order = [w["week"] for w in weeks]
    return {
        "weeks": weeks,
        "split": parse_split(wb),
        "sprint": parse_sprint(wb, order),
        "sollTotals": {
            key: round(sum(w["soll"].get(key, 0) for w in weeks))
            for key in ("at", "kraft", "spezKt", "wuerfe")
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", help="the plan spreadsheet")
    parser.add_argument("--key", help="service-account JSON")
    parser.add_argument("--name", help="plan name shown to athletes")
    parser.add_argument("--season", default="", help="e.g. 2026/27")
    parser.add_argument("--plan-id", help="overwrite this plan id")
    parser.add_argument("--dry-run", action="store_true",
                        help="parse and report, write nothing")
    parser.add_argument("--json-out", help="also write the parsed plan here")
    args = parser.parse_args()

    plan = build_plan(args.xlsx)
    season = args.season or _guess_season(plan)
    name = args.name or f"Trainingsplan {season}".strip()

    print(f"weeks       {len(plan['weeks'])} "
          f"({plan['weeks'][0]['week']} -> {plan['weeks'][-1]['week']})")
    print(f"soll totals {plan['sollTotals']}")
    print(f"split       {len(plan['split']['blocks'])} blocks, "
          f"{sum(len(b['exercises']) for b in plan['split']['blocks'])} "
          f"exercises")
    print(f"sprint      {len(plan['sprint']['blocks'])} blocks")

    problems = check(plan)
    if problems:
        print("\nPROBLEMS")
        for problem in problems:
            print(" ", problem)
        if not args.dry_run:
            return "refusing to upload; fix the sheet or re-run with --dry-run"

    if args.json_out:
        with open(args.json_out, "w", encoding="utf-8") as fh:
            json.dump(plan, fh, ensure_ascii=False, indent=2)
        print(f"\nwrote {args.json_out}")

    if args.dry_run:
        print("\ndry run, nothing uploaded")
        return 0

    if not args.key:
        return "--key is required to upload (or pass --dry-run)"
    if not os.path.exists(args.key):
        return f"no such service-account key: {args.key}"

    try:
        import firebase_admin
        from firebase_admin import credentials, firestore
    except ImportError:
        return "firebase-admin is missing:  pip install firebase-admin"

    firebase_admin.initialize_app(
        credentials.Certificate(args.key))
    db = firestore.client()

    owner = json.load(open(args.key, encoding="utf-8")).get("client_email", "")
    plan_id = args.plan_id or re.sub(r"[^A-Za-z0-9_-]", "-",
                                     f"plan-{season}").strip("-").lower()
    doc = db.collection("plans").document(plan_id)

    existing = doc.get()
    # Re-importing to fix a typo must not quietly unassign every athlete.
    assigned = (existing.to_dict() or {}).get("athleteUids", []) \
        if existing.exists else []
    owner_uid = (existing.to_dict() or {}).get("ownerUid") \
        if existing.exists else None

    if owner_uid is None:
        owner_uid = _prompt_owner_uid()

    payload = {
        **plan,
        "name": name,
        "season": season,
        "ownerUid": owner_uid,
        "athleteUids": assigned,
        "updatedAt": firestore.SERVER_TIMESTAMP,
        "importedBy": owner,
    }
    doc.set(payload)

    print(f"\nuploaded plans/{plan_id}")
    print(f"  name      {name}")
    print(f"  owner     {owner_uid}")
    print(f"  assigned  {len(assigned)} athlete(s) (unchanged)")
    print("\nAssign athletes in the app: Coach's Zone -> Pläne zuweisen.")
    return 0


def _guess_season(plan: dict) -> str:
    mondays = [w["monday"] for w in plan["weeks"] if w["monday"]]
    if not mondays:
        return ""
    start, end = mondays[0][:4], mondays[-1][:4]
    return start if start == end else f"{start}/{end[2:]}"


def _prompt_owner_uid() -> str:
    print("\nThis is a new plan. It needs the coach's Firebase Auth uid as "
          "owner —\nthe security rules only let that account edit or assign "
          "it.")
    print("Find it in the Firebase console under Authentication -> Users.")
    uid = input("Coach uid: ").strip()
    if not uid:
        sys.exit("no uid given")
    return uid


if __name__ == "__main__":
    sys.exit(main())
